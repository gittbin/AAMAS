"""One self-contained workbook for manual paper filling; no experiment changes.

Default export requires the complete study's successful raw-evidence audit.
--allow-partial produces an explicitly marked snapshot with missing cells blank.
Earlier studies remain in distinct sheets; smoke runs are never exported.
"""
from __future__ import annotations
import argparse
import csv
from datetime import datetime, timezone, timedelta
import hashlib
import json
import math
from pathlib import Path
import sys
import time

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from scripts_marl.run_paper_completion import ensure_gate, sha
from scripts_marl.report_paper_completion import matrix, METRICS
from scripts_marl.report_paper_v2 import nested_interval, exact_test, holm

MEASURES=['net_return','return_before_message_cost','return_at_default_price',
          'min_user_rate','sum_rate','illumination_db','outage_fraction',
          'requested','eligible','delivered','budget_dropped','graph_rejected',
          'reconfiguration','violation_before','violation_after','post25_return']
MEASURES += ['post25_'+m for m in MEASURES[3:15]]
BASELINES=('noop','random','heuristic','online_optimizer','ramp_knee','ramp_comm_extreme','ramp_sensing_extreme')
PROVENANCE=[]


def read(path):
    for attempt in range(5):
        try:
            value=json.loads(path.read_text(encoding='utf-8-sig'))
            PROVENANCE.append({'source_path':path.relative_to(ROOT).as_posix(),'sha256':sha(path)})
            return value
        except PermissionError:
            if attempt==4: raise
            time.sleep(0.1)


def category(condition,treatment):
    if condition=='baseline': return 'controller_reference'
    if treatment.startswith(('tau_','delay_','dynamics_','csi_')) or condition in ('feedforward','masked_csi'):
        return 'information_memory_robustness'
    if treatment.startswith('graph_') or condition=='selected_graph': return 'topology'
    if treatment.startswith(('frozen','static_')): return 'adaptation_static'
    if condition.startswith(('price_','budget3_')) or condition=='request_price': return 'price_budget'
    if condition in ('power','ris','ma','power_ris','power_ma','ris_ma'): return 'actuator'
    if condition=='full_comm' and treatment!='normal': return 'communication_intervention'
    return 'controller'


def average_rows(study,condition,seed,treatment,episodes,path,scope,audit_status,
                 checkpoint=None,step=None,steps=None):
    if not episodes: raise ValueError('No episodes: '+str(path))
    ids=[e['scenario_seed'] for e in episodes]
    if len(ids)!=len(set(ids)): raise ValueError('Duplicate scenarios: '+str(path))
    for e in episodes:
        if any(not math.isfinite(v) for v in e.values() if isinstance(v,(int,float))):
            raise ValueError('Nonfinite source episode: '+str(path))
    row={'study':study,'experiment':category(condition,treatment),'condition':condition,
         'treatment':treatment,'training_seed':seed,'n_scenarios':len(ids),
         'scenario_first':min(ids),'scenario_last':max(ids),'scope':scope,
         'audit_status':audit_status,'training_transitions':steps,'selected_step':step,
         'checkpoint_sha256':checkpoint,'source_json':path.relative_to(ROOT).as_posix()}
    for m in MEASURES:
        values=[e[m] for e in episodes if m in e]
        row[m]=float(np.mean(values)) if len(values)==len(episodes) else None
    row['return_sd_across_scenarios']=float(np.std([e['net_return'] for e in episodes],ddof=1)) if len(episodes)>1 else None
    return row


def expected_treatments(condition,p):
    expected={'normal'}
    if condition=='full_nocomm': expected|={'frozen25','frozen10','static_learned'}|{'csi_'+m for m in p['csi_feature_interventions']}
    if condition=='full_comm':
        expected|={'comm_off','comm_random','comm_forced','content_zero','content_random','content_permute'}
        expected|={'graph_'+g for g in ('empty','full','ring','bs_star','selected','harmful')}
        expected|={'graph_harmful_content_'+g for g in ('zero','random','permute')}
    if condition in ('full_nocomm','feedforward','masked_csi'):
        expected|={f'tau_{v}' for v in p['tau_grid']}|{f'delay_{v}' for v in p['delay_grid']}|{'dynamics_'+m for m in ('static','slow','fast')}
        if condition=='full_nocomm': expected|={'dynamics_'+m+'_frozen25' for m in ('static','slow','fast')}
    return expected


def gather_main(out,allow_partial):
    lock=read(out/'execution_protocol.json'); manifest=ensure_gate(out,lock); p=lock['protocol']
    evidence=read(out/'report/audit_final.json') if (out/'report/audit_final.json').exists() else {}
    final_report=read(out/'report/results.json') if (out/'report/results.json').exists() else {}
    full_audit=evidence.get('audit_passed',False) and final_report.get('audit_passed',False)
    if not full_audit and not allow_partial: raise RuntimeError('Full evidence audit is still pending; wait or explicitly use --allow-partial')
    correction=read(out/'numerical_erratum_locked.json')
    for name,digest in correction['correction_source_files'].items():
        if sha(ROOT/name)!=digest: raise ValueError('Numerical correction source changed')
    if full_audit and sha(out/'numerical_erratum_locked.json')!=final_report['numerical_erratum']['correction_lock_sha256']:
        raise ValueError('Correction ledger changed since audit')
    status='FULL_CAMPAIGN_AUDIT_PASSED' if full_audit else 'EPISODE_AND_CHECKPOINT_CHECKED_FULL_AUDIT_PENDING'
    dataset,per_seed,training={},[],[]
    ids=list(range(p['reserved_final_seed_start'],p['reserved_final_seed_start']+p['reserved_final_scenarios']))
    for condition in p['conditions']:
        for seed in p['training_seeds']:
            train=out/f'train/{condition}/seed_{seed}'; s=read(train/'status.json')
            if not s['training_complete'] or s['actual_transitions']!=p['steps_per_run']:
                raise ValueError('Incomplete training')
            entry=manifest['selected'][f'{condition}/seed_{seed}']
            if sha(train/'checkpoint_best.pt')!=entry['selected_checkpoint_sha256']:
                raise ValueError('Checkpoint mismatch')
            path=out/f'evaluation/{condition}/seed_{seed}/results.json'; d=read(path)
            if not d['complete'] or d['checkpoint_lock']!=entry or d['source_sha256']!=lock['source_sha256']:
                raise ValueError('Invalid final policy provenance')
            if set(d['treatments'])!=expected_treatments(condition,p) or d['scenario_ids']!=ids:
                raise ValueError('Missing treatment or mismatched final scenario IDs')
            config=s['config']; communication=config['communication']
            training.append({'condition':condition,'training_seed':seed,'training_transitions':s['actual_transitions'],
                'selected_step':entry['selected_step'],'selected_development_return':s['best_development_return'],
                'actor_parameters':s['policy_parameters'],'critic_parameters':s['critic_parameters'],
                'delivery_budget':communication['max_messages_per_step'],'message_price':config['reward']['w_message'],
                'selected_checkpoint_sha256':entry['selected_checkpoint_sha256'],'final_checkpoint_sha256':entry['final_checkpoint_sha256'],
                'selection_log_sha256':entry['selection_log_sha256'],'initial_actor_sha256':s['initial_actor_sha256'],
                'source_sha256':s['source_sha256'],'source_status':(train/'status.json').relative_to(ROOT).as_posix()})
            for treatment,item in d['treatments'].items():
                eps=item['episodes']; expected=ids[:p['exploratory_final_scenarios']] if treatment.startswith(('tau_','delay_','dynamics_')) else ids
                if [e['scenario_seed'] for e in eps]!=expected: raise ValueError('Scenario coverage mismatch')
                scope='fresh_final_exploratory' if treatment.startswith(('tau_','delay_','dynamics_','csi_','graph_','content_','comm_random','comm_forced','frozen10','static_')) or condition.startswith(('price_','budget3_')) or condition in ('request_price','feedforward','masked_csi','selected_graph','centralized') else 'fresh_final_prespecified_or_nominal'
                row=average_rows('fresh_complete_10_seed',condition,seed,treatment,eps,path,scope,status,
                    entry['selected_checkpoint_sha256'],entry['selected_step'],s['actual_transitions'])
                row.update(delivery_budget=communication['max_messages_per_step'],message_price=config['reward']['w_message'],
                    numerical_erratum='perfect_CSI_overflow_only' if treatment=='tau_0.0' else 'none_original_evaluator')
                per_seed.append(row); dataset.setdefault(condition+'/'+treatment,[]).append((seed,eps))
    base=read(out/'baselines/results.json') if (out/'baselines/results.json').exists() else {'complete':False,'references':{}}
    reference_rows=[]
    for name,item in base['references'].items():
        eps=item['episodes']
        if [e['scenario_seed'] for e in eps]!=ids: raise ValueError('Baseline coverage mismatch')
        dataset['baseline/'+name]=[(None,eps)]
        reference_rows.append(average_rows('fresh_complete_10_seed','baseline',None,name,eps,out/'baselines/results.json',
            'nonlearning_reference_no_training_replicates',status))
    if full_audit and (not base['complete'] or set(base['references'])!=set(BASELINES)):
        raise ValueError('Full audit but incomplete baseline snapshot')
    # When an audited report exists, every exported seed mean must agree with it.
    if full_audit:
        official={(r['condition'],r['treatment'],r['seed']):r for r in final_report['per_seed']}
        for row in per_seed+reference_rows:
            key=(row['condition'],row['treatment'],row['training_seed'])
            if key not in official: raise ValueError('Exported cell absent from audited report')
            for metric in METRICS:
                if not math.isclose(row[metric],official[key][metric],rel_tol=1e-10,abs_tol=1e-10):
                    raise ValueError('Seed mean disagrees with audited report')
    return p,dataset,per_seed,reference_rows,training,evidence,final_report,full_audit


def comparisons(dataset,p,final_report):
    primary,seed_differences=[],[]
    for h in p['primary_hypotheses']:
        a,b=matrix(dataset,h['left'],h['metric']),matrix(dataset,h['right'],h['metric'])
        effect=nested_interval(a,b,paired=True,repetitions=p['bootstrap_repetitions'],seed=p['statistics_seed'])
        differences=a.mean(1)-b.mean(1)
        primary.append(dict(h,mean=effect['mean'],ci95_low=effect['ci95'][0],ci95_high=effect['ci95'][1],
            p_exact=exact_test(a,b,paired=True),positive_seed_pairs=int((differences>0).sum()),training_pairs=len(a),scenarios_per_policy=a.shape[1]))
        for i,seed in enumerate(p['training_seeds']):
            seed_differences.append({'comparison':h['name'],'scope':'primary_Holm_family_7','left':h['left'],'right':h['right'],
                'metric':h['metric'],'training_seed':seed,'left_mean':float(a[i].mean()),'right_mean':float(b[i].mean()),'left_minus_right':float(differences[i])})
    for item,adjusted in zip(primary,holm([h['p_exact'] for h in primary])):
        item['p_holm']=adjusted; item['supported_at_Holm_0.05']=adjusted<0.05
    if final_report:
        for row,official in zip(primary,final_report['primary']):
            for key in ('mean','p_exact','p_holm'):
                if not math.isclose(row[key],official[key],rel_tol=1e-10,abs_tol=1e-10): raise ValueError('Primary test mismatch')
            if not np.allclose([row['ci95_low'],row['ci95_high']],official['ci95'],rtol=1e-10,atol=1e-10): raise ValueError('Primary interval mismatch')
    exploratory=[]
    specs=[('content_'+m,'full_comm/normal','full_comm/content_'+m,'net_return') for m in ('zero','random','permute')]
    specs += [('recurrent_vs_feedforward','full_nocomm/normal','feedforward/normal','net_return'),
              ('recurrent_vs_masked','full_nocomm/normal','masked_csi/normal','net_return'),
              ('harmful_content_zero','full_comm/graph_harmful','full_comm/graph_harmful_content_zero','net_return')]
    specs += [('adaptation_'+v,'full_nocomm/dynamics_'+v,'full_nocomm/dynamics_'+v+'_frozen25','post25_min_user_rate') for v in ('static','slow','fast')]
    for name,left,right,metric in specs:
        a,b=matrix(dataset,left,metric),matrix(dataset,right,metric)
        effect=nested_interval(a,b,paired=True,repetitions=p['bootstrap_repetitions'],seed=p['statistics_seed'])
        exploratory.append({'comparison':name,'left':left,'right':right,'metric':metric,'mean':effect['mean'],
            'ci95_low':effect['ci95'][0],'ci95_high':effect['ci95'][1],'scope':'exploratory_no_primary_adjusted_test'})
        for i,seed in enumerate(p['training_seeds']):
            seed_differences.append({'comparison':name,'scope':'exploratory','left':left,'right':right,'metric':metric,
                'training_seed':seed,'left_mean':float(a[i].mean()),'right_mean':float(b[i].mean()),'left_minus_right':float(a[i].mean()-b[i].mean())})
    return primary,exploratory,seed_differences


def summary_rows(dataset,p,final_report):
    results=[]
    for key in sorted(dataset):
        condition,treatment=key.split('/'); entries=dataset[key]
        row={'experiment':category(condition,treatment),'condition':condition,'treatment':treatment,
             'n_training_seeds':0 if condition=='baseline' else len(entries),'n_scenarios_per_policy':len(entries[0][1]),
             'training_steps_per_seed':None if condition=='baseline' else p['steps_per_run'],
             'status':'AVAILABLE','scope':'nonlearning_no_train_seed' if condition=='baseline' else 'fresh_final'}
        for metric in MEASURES:
            values=matrix(dataset,key,metric)
            row[metric+'_mean']=float(values.mean())
            row[metric+'_sd_between_training_seeds']=float(values.mean(1).std(ddof=1)) if len(values)>1 else None
        for metric in ('net_return','min_user_rate','illumination_db'):
            estimate=final_report.get('summaries',{}).get(key,{}).get('metrics',{}).get(metric,{})
            # Use the audited report interval where available; other treatment
            # intervals are pointwise descriptive estimates under the same locked
            # bootstrap scheme, without adding or reselecting primary tests.
            ci=estimate.get('ci95')
            if ci is None:
                ci=nested_interval(matrix(dataset,key,metric),repetitions=p['bootstrap_repetitions'],seed=p['statistics_seed'])['ci95']
            row[metric+'_ci95_low'],row[metric+'_ci95_high']=ci
        results.append(row)
    for name in BASELINES:
        if 'baseline/'+name not in dataset:
            results.append({'experiment':'controller_reference','condition':'baseline','treatment':name,
                'status':'PENDING_NOT_MEASURED_YET','scope':'no fabricated zeros'})
    return results


def original_template_rows(summaries,training,timing,graph,p):
    """Map results to the author's original results_template.csv column names."""
    lookup={r['condition']+'/'+r['treatment']:r for r in summaries}
    setup={r['condition']:r for r in training if r['training_seed']==p['training_seeds'][0]}
    mappings=[('controller','uspa','baseline/noop'),('controller','archive_knee','baseline/ramp_knee'),
        ('controller','archive_comm_extreme','baseline/ramp_comm_extreme'),('controller','archive_sensing_extreme','baseline/ramp_sensing_extreme'),
        ('controller','heuristic','baseline/heuristic'),('controller','random','baseline/random'),
        ('controller','policy_frozen','full_nocomm/frozen25'),('controller','policy_selected_static','full_nocomm/static_learned'),
        ('controller','mappo_train_nocomm','full_nocomm/normal'),('controller','mappo_train_comm','full_comm/normal'),
        ('controller','centralized_ppo','centralized/normal'),('controller','online_optimizer','baseline/online_optimizer')]
    mappings += [('actuator',c,c+'/normal') for c in ('power','ris','ma','power_ris','power_ma','ris_ma')]
    mappings += [('actuator','full','full_nocomm/normal')]
    mappings += [('communication',alias,key) for alias,key in (
        ('train_nocomm','full_nocomm/normal'),('normal','full_comm/normal'),('eval_off','full_comm/comm_off'),
        ('zero_fixed_delivery','full_comm/content_zero'),('random_fixed_delivery','full_comm/content_random'),
        ('permute_fixed_delivery','full_comm/content_permute'),('delivery_random','full_comm/comm_random'),('delivery_forced','full_comm/comm_forced'))]
    mappings += [('price',c,c+'/normal') for c in ('price_free','full_comm','price_tenfold','budget3_free','budget3_default','budget3_tenfold','request_price')]
    mappings += [('graph',g,'full_comm/graph_'+g) for g in ('empty','full','ring','bs_star','selected','harmful')]
    mappings += [('graph','selected_graph_retrained','selected_graph/normal')]
    mappings += [('memory',alias,key) for alias,key in (('recurrent','full_nocomm/normal'),('feedforward','feedforward/normal'),
        ('masked_csi_trained','masked_csi/normal'),('fixed_csi_eval','full_nocomm/csi_fixed'),('shuffled_csi_eval','full_nocomm/csi_shuffled'))]
    mappings += [('robustness',key,key) for key in lookup if any('/'+m in key for m in ('tau_','delay_','dynamics_'))]
    with (ROOT/'AAMAS_rewrite_complete (1)/results_template.csv').open(newline='',encoding='utf-8-sig') as f:
        columns=next(csv.reader(f))
    results=[]
    for experiment,alias,key in mappings:
        if key not in lookup: continue
        data=lookup[key]; condition,treatment=key.split('/'); cfg=setup.get(condition,{})
        row=dict.fromkeys(columns)
        row.update(experiment=experiment,condition=alias,tau=0.1,actor_delay=1,velocity='0.5–2 m/s',rho=0.95,
            budget=cfg.get('delivery_budget',0),message_price=cfg.get('message_price',0),
            training_steps=data.get('training_steps_per_seed'),training_seeds=data.get('n_training_seeds'),
            evaluation_scenarios=data.get('n_scenarios_per_policy'),
            notes='Source '+key+'; outcome/test interpretation in PRIMARY_7. CI95 pointwise; scenarios are not training seeds.')
        if condition=='baseline': row['training_seeds']=0
        if treatment.startswith('tau_'): row['tau']=float(treatment[4:])
        if treatment.startswith('delay_'): row['actor_delay']=int(treatment[6:])
        if treatment.startswith('dynamics_'):
            mode=treatment[len('dynamics_'):].split('_')[0]
            row['velocity'],row['rho']={'static':('0 m/s',1.0),'slow':('0–0.5 m/s',0.99),'fast':('5–10 m/s',0.7)}[mode]
        if treatment.startswith('graph_'): row['graph']=graph['graphs_for_final'].get(treatment[6:])
        if condition=='selected_graph': row['graph']=graph['selected_graph']
        if 'frozen25' in treatment: row['freeze_time']=25
        if treatment=='tau_0.0': row['notes']+=' Disclosed post-final perfect-CSI overflow erratum.'
        conversion={'return':'net_return','min_rate':'min_user_rate','sum_rate':'sum_rate',
            'illumination':'illumination_db','outage':'outage_fraction','reconfiguration':'reconfiguration',
            'requests':'requested','eligible':'eligible','delivered':'delivered','drops':'budget_dropped'}
        for target,source in conversion.items():
            row[target+'_mean']=data.get(source+'_mean')
            for suffix in ('ci_low','ci_high'):
                if target+'_'+suffix in row:
                    row[target+'_'+suffix]=data.get(source+'_ci95_'+suffix[3:])
        measured='online_optimizer' if condition=='baseline' and treatment=='online_optimizer' else condition
        if treatment=='normal' or key=='baseline/online_optimizer':
            if measured in timing:
                row['latency_median_ms']=timing[measured]['summary']['decision_kernels_ms']['median']
                row['latency_p95_ms']=timing[measured]['summary']['decision_kernels_ms']['p95']
        row.update(source_condition_treatment=key,data_status=data['status'],
            return_sd_between_training_seeds=data.get('net_return_sd_between_training_seeds'),
            min_rate_sd_between_training_seeds=data.get('min_user_rate_sd_between_training_seeds'))
        results.append(row)
    return results


def historical():
    outputs={}
    studies=[('V3_CONFIRM','paper_v3/confirmation_v2','separate_final_8_seeds_do_not_pool','audit_final.json'),
             ('V4_SCREEN','paper_v4/screening','development_screen_3_seeds_not_final','audit.json'),
             ('V4_DIAG','paper_v4/diagnostics','development_diagnostic_not_final','audit.json'),
             ('V4_ADAPT','paper_v4/adaptation_development','development_adaptation_not_final','audit.json'),
             ('V3_SCREEN','paper_v3/screening','development_screen_2_seeds_not_final',None),
             ('V2_LEGACY','paper_v2/campaign','DIAGNOSTIC_ONLY_TRAIN_FINAL_SEED_OVERLAP','audit.json')]
    for sheet,study,scope,audit_name in studies:
        folder=ROOT/'outputs_marl'/study
        evidence=read(folder/'report'/audit_name) if audit_name and (folder/'report'/audit_name).exists() else {}
        audit_state='PRIOR_STUDY_AUDIT_PASSED' if evidence.get('audit_passed') else 'SOURCE_SUMMARY_ONLY_NO_NEW_RAW_AUDIT'
        if sheet=='V2_LEGACY': audit_state='NUMERIC_AUDIT_DOES_NOT_REMOVE_SEED_LEAKAGE'
        records=[]
        files=sorted((folder/'evaluation').glob('*/seed_*/results.json')) or sorted(folder.glob('seed_*/results.json'))
        for path in files:
            d=read(path)
            if not d.get('complete'): raise ValueError('Incomplete historical record: '+str(path))
            condition=d.get('condition','no_comm' if sheet=='V4_ADAPT' else 'comm_reference')
            status_path=folder/f'train/{condition}/seed_{d["training_seed"]}/status.json'
            status=read(status_path) if status_path.exists() else {}
            for treatment,item in d['treatments'].items():
                eps=item if isinstance(item,list) else item['episodes']
                entry=d.get('checkpoint_lock',{})
                final=treatment in ('final','final_checkpoint')
                checkpoint=(d.get('checkpoint_final_sha256') or entry.get('final_checkpoint_sha256')) if final else (
                    d.get('checkpoint_sha256') or d.get('checkpoint_best_sha256') or entry.get('selected_checkpoint_sha256'))
                steps=status.get('actual_transitions')
                step=steps if final else entry.get('selected_step',entry.get('best_step',status.get('best_step')))
                records.append(average_rows(study,condition,d['training_seed'],treatment,eps,path,scope,audit_state,checkpoint,step,steps))
        outputs[sheet]=records
    folder=ROOT/'outputs_marl/paper_v2/diagnostics'; path=folder/'results.json'; d=read(path); capability=[]
    for run in d['runs']:
        for phase in ('before','after'):
            capability.append({'task':run['task'],'price':run['price'],'training_seed':run['training_seed'],'phase':phase,
                'training_transitions':0 if phase=='before' else run['actual_transitions'],
                **run[phase],'diagnostics_passed':d['passed'],'scope':'toy_diagnostic_not_wireless_performance',
                'source_json':path.relative_to(ROOT).as_posix()})
    outputs['COMM_CAPABILITY']=capability
    return outputs


def add_sheet(book,name,rows):
    ws=book.create_sheet(name)
    if not rows: ws.append(['status']); ws.append(['PENDING / NO DATA']); return ws
    columns=list(dict.fromkeys(key for row in rows for key in row))
    ws.append(columns)
    for row in rows:
        values=[]
        for c in columns:
            value=row.get(c)
            if isinstance(value,(dict,list)): value=json.dumps(value,ensure_ascii=False)
            if isinstance(value,str) and value.startswith(('=','+','-','@')): value="'"+value
            values.append(value)
        ws.append(values)
    ws.freeze_panes='D2'; ws.sheet_view.zoomScale=85
    table=Table(displayName='T_'+name,ref=f'A1:{get_column_letter(len(columns))}{len(rows)+1}')
    table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
    ws.add_table(table)
    for cell in ws[1]:
        cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='183D5C')
        cell.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=46
    for i,c in enumerate(columns,1):
        ws.column_dimensions[get_column_letter(i)].width=min(44,max(15,len(c)*0.65))
        for column_cell in ws.iter_cols(min_col=i,max_col=i,min_row=2):
            for cell in column_cell:
                if isinstance(cell.value,float): cell.number_format='0.000000'
                elif isinstance(cell.value,int): cell.number_format='0'
                elif c in ('note','notes','scope','status','audit_status','value'):
                    cell.alignment=Alignment(wrap_text=True,vertical='top')
    return ws


def export(out,destination,allow_partial=False):
    p,dataset,seed_rows,refs,training,evidence,official,full_audit=gather_main(out,allow_partial)
    primary,exploratory,differences=comparisons(dataset,p,official)
    summaries=summary_rows(dataset,p,official)
    histories=historical(); graph=read(out/'graph_selection.json'); graph_rows=[]
    for seed in p['training_seeds']:
        path=out/f'graph_development/seed_{seed}/results.json'; d=read(path)
        for g,eps in d['episodes_by_graph'].items():
            row=average_rows('fresh_graph_development','full_comm',seed,'graph_'+g,eps,path,
                'DEVELOPMENT_SELECTION_NOT_FINAL','FULL_SELECTION_AUDIT_PASSED' if full_audit else 'SELECTION_AUDIT_PENDING',d['checkpoint_sha256'])
            row.update(graph_integer=int(g),selected_for_retraining=int(g)==graph['selected_graph'],
                       selected_harmful=int(g)==graph['harmful_graph'])
            graph_rows.append(row)
    timestamp=datetime.now(timezone(timedelta(hours=7))).strftime('%d/%m/%Y %H:%M:%S UTC+7')
    missing=[name for name in BASELINES if 'baseline/'+name not in dataset]
    timing=read(out/'baselines/latency.json') if (out/'baselines/latency.json').exists() else {}
    if not timing: missing.append('latency')
    if not full_audit: missing.append('full_campaign_audit')
    ready=not missing and full_audit
    book=openpyxl.Workbook(); book.remove(book.active)
    guide=[{'item':'Thời điểm xuất','value':timestamp},
        {'item':'Đủ dữ liệu điền toàn bộ phần Results?','value':'CÓ: đủ dữ liệu đã audit; tác giả tự điền và diễn giải.' if ready else 'CHƯA đủ toàn bộ. Có kết quả từng seed cho các nhánh học; các mục thiếu được ghi rõ.'},
        {'item':'Mục còn thiếu','value':', '.join(missing) if missing else 'Không thiếu dữ liệu thí nghiệm trong protocol đã khóa.'},
        {'item':'TONG_HOP_PAPER','value':'Mean, sample SD giữa training seed và CI95 cho return/min-rate/illumination; mỗi hàng là một condition/treatment. Dùng CI đã audit nếu có; các treatment khác tính CI mô tả pointwise từ episode đã audit theo cùng bootstrap đã khóa.'},
        {'item':'BANG_DIEN_GOC','value':'Cột theo results_template.csv của paper gốc: return_mean, min_rate_mean, CI95, message budget/price, graph, freeze, latency. Các kiểm định chính ở PRIMARY_7; ô trống không được thay bằng số 0.'},
        {'item':'TUNG_SEED_FINAL','value':'770 hàng: toàn bộ 18 nhánh × 10 seed và các can thiệp. Mỗi giá trị là trung bình episode của một policy seed trên 150 scenario chung; robustness dùng 30 scenario.'},
        {'item':'CHENH_LECH_SEED / PRIMARY_7','value':'Hiệu ứng LEFT MINUS RIGHT từng seed; 7 kiểm định chính, CI pointwise và p exact/Holm của cả family.'},
        {'item':'BASELINES','value':'Các reference không train không có training_seed. Một reference dùng 150 scenario; không tạo 10 seed giả.'},
        {'item':'GRAPH_DEVELOPMENT','value':'640 hàng: 64 graph × 10 checkpoint trên 10 scenario development. Không dùng làm bảng final.'},
        {'item':'CHECKPOINTS','value':'180 lượt train, 400384 transitions mỗi lượt, best-step và SHA256; development return dùng chọn checkpoint, không phải final return.'},
        {'item':'Nghiên cứu cũ','value':'V3_CONFIRM là final riêng với 8 seed; V3_SCREEN/V4_SCREEN/V4_DIAG/V4_ADAPT là development; V2_LEGACY chỉ diagnostic vì seed leakage. Không gộp các study.'},
        {'item':'net_return','value':'Tổng reward 100 bước/episode, có message price theo từng condition. return_before_message_cost và return_at_default_price tách ảnh hưởng giá.'},
        {'item':'min_user_rate / sum_rate','value':'Trung bình theo bước, đơn vị bit/s/Hz. post25_* chỉ lấy bước 25–99; post25_return là tổng reward 75 bước.'},
        {'item':'illumination_db / outage_fraction','value':'Illumination tính bằng dB theo convention evaluator. Outage là phần người dùng dưới 1 bit/s/Hz, trung bình theo bước; không phải xác suất cả episode outage.'},
        {'item':'requested / eligible / delivered / budget_dropped / graph_rejected','value':'Số tin mỗi control interval. reconfiguration và violations theo thang simulator.'},
        {'item':'Phân biệt SD','value':'*_sd_between_training_seeds: sample SD của policy-seed means; return_sd_across_scenarios: SD episode trong một seed. Scenario không được tính là training replicate.'},
        {'item':'CI và p','value':'Nested bootstrap 10000 lần trên training blocks và shared scenario IDs; exact sign flip hai phía trên 10 block means; Holm chỉnh cả 7 primary tests. CI là pointwise; nonsignificance không chứng minh equivalence.'},
        {'item':'Numerical erratum sau khi mở final','value':official.get('numerical_erratum',{}).get('disclosure','Đã sửa overflow inverse-squared power normalization chỉ ở tau=0.0; chạy lại đồng nhất 30 policy cells. Primary và nominal giữ evaluator/trace gốc; không train hoặc chọn checkpoint/seed lại.')},
        {'item':'Diễn giải có bằng chứng','value':'Adaptation sau freeze25 cải thiện rate; full hơn MA, power+RIS và RIS+MA. Chưa xác nhận full hơn power+MA hoặc lợi ích communication. Xem CI và Holm trong PRIMARY_7.'},
        {'item':'Scope của timing','value':'CPU một Torch thread, batch 1, decision kernels; không phải end-to-end deployment latency.'},
        {'item':'File này có sửa paper/source/thí nghiệm không?','value':'Không. Đây là bản xuất để tác giả tự điền. Dữ liệu smoke, incomplete evaluation và aborted original tau=0 không được gộp vào kết quả.'}]
    add_sheet(book,'HUONG_DAN',guide)
    add_sheet(book,'BANG_DIEN_GOC',original_template_rows(summaries,training,timing,graph,p))
    add_sheet(book,'TONG_HOP_PAPER',summaries)
    add_sheet(book,'TUNG_SEED_FINAL',seed_rows)
    add_sheet(book,'CHENH_LECH_SEED',differences)
    add_sheet(book,'PRIMARY_7',primary)
    add_sheet(book,'EXPLORATORY_CI',exploratory)
    add_sheet(book,'BASELINES',refs+[{'condition':'baseline','treatment':name,'audit_status':'PENDING','scope':'not measured; all metric cells blank'} for name in BASELINES if 'baseline/'+name not in dataset])
    timing_rows=[]; timing_samples=[]
    for condition,d in timing.items():
        for component,summary in d['summary'].items():
            timing_rows.append({'condition':condition,'component':component,**summary,'n_measurements':d['n'],
                'benchmark_policy_training_seed':None if condition=='online_optimizer' else p['training_seeds'][0],
                'warmup':d['warmup'],'torch_threads':d['torch_threads'],'scope':d['scope']})
        for index in range(d['n']):
            timing_samples.append({'condition':condition,'sample_index':index,
                'benchmark_policy_training_seed':None if condition=='online_optimizer' else p['training_seeds'][0],
                **{component:values[index] for component,values in d['samples'].items()}})
    add_sheet(book,'LATENCY',timing_rows)
    add_sheet(book,'LATENCY_SAMPLES',timing_samples)
    add_sheet(book,'CHECKPOINTS',training)
    add_sheet(book,'GRAPH_DEVELOPMENT',graph_rows)
    for name,rows in histories.items(): add_sheet(book,name,rows)
    checklist=[]
    for group in ('controller','adaptation_static','actuator','information_memory_robustness','communication_intervention','price_budget','topology','controller_reference'):
        keys=[k for k in dataset if category(*k.split('/'))==group]
        checklist.append({'paper_results_group':group,'available_condition_treatments':len(keys),
            'keys':', '.join(keys),'status':'AUDITED_AVAILABLE' if full_audit else 'AVAILABLE_FULL_AUDIT_PENDING'})
    checklist.append({'paper_results_group':'latency','status':'AUDITED_AVAILABLE' if full_audit and timing else 'PENDING'})
    add_sheet(book,'MAP_PAPER',checklist)
    provenance=list({r['source_path']:r for r in PROVENANCE}.values())
    provenance.append({'source_path':'scripts_marl/export_seed_workbook.py','sha256':sha(Path(__file__).resolve())})
    add_sheet(book,'SOURCES_SHA256',provenance)
    ws=book['HUONG_DAN']; ws.column_dimensions['A'].width=44; ws.column_dimensions['B'].width=112; ws.freeze_panes='A2'
    for i in range(2,len(guide)+2): ws.row_dimensions[i].height=44
    book.properties.title='AAMAS: audited results per training seed and experiment'
    book.properties.creator='Research project'; book.properties.description='Separate studies, scopes, immutable selected policies, paired inference, disclosed numerical erratum.'
    destination.parent.mkdir(parents=True,exist_ok=True)
    if destination.exists():
        archived=destination.with_name(destination.stem+'_previous_'+datetime.now().strftime('%Y%m%d_%H%M%S')+destination.suffix)
        destination.replace(archived)
    book.save(destination)
    # Verify numeric cells, complete coverage, no unexpected Excel formulas and
    # successful serialization, without changing frozen experiment artifacts.
    loaded=openpyxl.load_workbook(destination,read_only=True,data_only=False)
    if loaded['TUNG_SEED_FINAL'].max_row-1!=770 or loaded['CHECKPOINTS'].max_row-1!=180:
        raise ValueError('Exported coverage is incomplete')
    if loaded['PRIMARY_7'].max_row-1!=7: raise ValueError('Primary family size changed')
    for ws in loaded:
        for row in ws:
            for cell in row:
                if cell.data_type=='f': raise ValueError('Unexpected spreadsheet formula')
                if isinstance(cell.value,float) and not math.isfinite(cell.value): raise ValueError('Nonfinite workbook cell')
    sheets={ws.title:ws.max_row-1 for ws in loaded}; loaded.close()
    print(json.dumps({'path':str(destination),'bytes':destination.stat().st_size,'ready_for_results_filling':ready,
        'full_audit_passed':full_audit,'missing':missing,'sheets':sheets,'verification_passed':True},indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--campaign',type=Path,default=ROOT/'outputs_marl/paper_completion_recovery')
    parser.add_argument('--out',type=Path); parser.add_argument('--allow-partial',action='store_true')
    args=parser.parse_args(); campaign=args.campaign.resolve()
    export(campaign,args.out.resolve() if args.out else campaign/'report/KET_QUA_TUNG_SEED_DIEN_PAPER.xlsx',args.allow_partial)
